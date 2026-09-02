"""Tests for Excel (.xlsx) multi-tab Job P&L export and spreadsheet sync."""

import pytest
import openpyxl
from pathlib import Path
from invoiceledger.models import Job, ExtractedInvoice, InvoiceStatus, ColumnMappingTemplate
from invoiceledger.spreadsheet import SpreadsheetManager


def test_export_to_excel_multi_tab(tmp_path):
    jobs = [
        Job(
            id="job-101",
            name="101 Main St Remodel",
            client_name="Oakwood Holdings",
            budget_total=85000.0,
            cost_code_budgets={"06-Framing": 25000.0, "07-Roofing": 15000.0},
        )
    ]

    invoices = [
        ExtractedInvoice(
            id="inv-1",
            source_filename="hd.pdf",
            vendor_name="The Home Depot",
            invoice_number="HD-101",
            invoice_date="2026-08-10",
            total_amount=1450.75,
            job_name="101 Main St Remodel",
            cost_code="06-Framing",
            status=InvoiceStatus.COMMITTED,
            file_hash="h1",
            archive_path="archive/101-main/hd_101.pdf",
        )
    ]

    out_xlsx = tmp_path / "Contractor_PL.xlsx"
    raw_bytes, summary = SpreadsheetManager.export_to_excel(invoices, jobs, output_path=out_xlsx)

    assert out_xlsx.exists()
    assert summary.total_amount == 1450.75
    assert summary.is_reconciled is True

    # Load workbook and verify sheets
    wb = openpyxl.load_workbook(out_xlsx)
    assert "Job P&L Summary" in wb.sheetnames
    assert "Cost Code Breakdown" in wb.sheetnames
    assert "Invoice Ledger" in wb.sheetnames

    ws_summary = wb["Job P&L Summary"]
    assert ws_summary["B4"].value == "101 Main St Remodel"
    assert ws_summary["D4"].value == 85000.0

    ws_ledger = wb["Invoice Ledger"]
    assert ws_ledger["D4"].value == "The Home Depot"
    assert ws_ledger["G4"].value == 1450.75


def test_update_job_pl_spreadsheet_csv(tmp_path):
    job = Job(
        id="job-204",
        name="204 Pine Ridge Commercial",
        client_name="Apex Corp",
        budget_total=175000.0,
        cost_code_budgets={"15-Mechanical/Plumbing": 40000.0},
    )

    invoices = [
        ExtractedInvoice(
            id="inv-ferg",
            source_filename="ferg.pdf",
            vendor_name="Ferguson Enterprises",
            invoice_number="FERG-10940",
            invoice_date="2026-08-04",
            total_amount=3640.50,
            job_name="204 Pine Ridge Commercial",
            cost_code="15-Mechanical/Plumbing",
            status=InvoiceStatus.COMMITTED,
            file_hash="h_ferg",
        )
    ]

    target_csv = tmp_path / "Pine_Ridge_PL.csv"
    ok = SpreadsheetManager.update_job_pl_spreadsheet(job, invoices, target_csv)
    assert ok is True
    assert target_csv.exists()
    content = target_csv.read_text(encoding="utf-8-sig")
    assert "Ferguson Enterprises" in content
    assert "3640.50" in content


def test_update_job_pl_existing_workbook_in_place(tmp_path):
    """Verifies that updating an existing workbook appends new rows without destroying other sheets."""
    job = Job(
        id="job-101",
        name="101 Main St Remodel",
        client_name="Oakwood Holdings",
        budget_total=85000.0,
        cost_code_budgets={"06-Framing": 25000.0},
    )

    existing_xlsx = tmp_path / "Existing_Job_PL.xlsx"

    # Create initial workbook with invoice 1
    inv1 = ExtractedInvoice(
        id="inv-1",
        source_filename="hd1.pdf",
        vendor_name="The Home Depot",
        invoice_number="HD-101",
        invoice_date="2026-08-10",
        total_amount=1000.00,
        job_name="101 Main St Remodel",
        cost_code="06-Framing",
        status=InvoiceStatus.COMMITTED,
        file_hash="h1",
    )
    SpreadsheetManager.update_job_pl_spreadsheet(job, [inv1], existing_xlsx)

    # Now append invoice 2 to the same workbook
    inv2 = ExtractedInvoice(
        id="inv-2",
        source_filename="hd2.pdf",
        vendor_name="ABC Supply Co",
        invoice_number="ABC-202",
        invoice_date="2026-08-12",
        total_amount=2500.00,
        job_name="101 Main St Remodel",
        cost_code="06-Framing",
        status=InvoiceStatus.COMMITTED,
        file_hash="h2",
    )
    ok = SpreadsheetManager.update_job_pl_spreadsheet(job, [inv1, inv2], existing_xlsx)
    assert ok is True

    # Check updated workbook
    wb = openpyxl.load_workbook(existing_xlsx)
    ws_ledger = wb["Invoice Ledger"]
    inv_nums_in_sheet = [ws_ledger.cell(row=r, column=5).value for r in range(4, ws_ledger.max_row + 1)]
    assert "HD-101" in inv_nums_in_sheet
    assert "ABC-202" in inv_nums_in_sheet
    assert inv_nums_in_sheet.count("HD-101") == 1  # No duplicates
