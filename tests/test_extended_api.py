"""Tests for extended API endpoints (vendors, Excel export, spreadsheet sync, scheduler)."""

import pytest
from fastapi.testclient import TestClient
from pathlib import Path
from invoiceledger.server import app, storage

client = TestClient(app)


def test_api_vendors_crud():
    res = client.get("/api/vendors")
    assert res.status_code == 200
    vendors = res.json()
    assert len(vendors) >= 5

    new_vendor = {
        "id": "v-test-api",
        "name": "Acme Drywall Supply",
        "aliases": ["Acme Drywall", "Acme Supply"],
        "default_job_name": "101 Main St Remodel",
        "default_cost_code": "09-Finishes",
    }
    create_res = client.post("/api/vendors", json=new_vendor)
    assert create_res.status_code == 200
    assert create_res.json()["name"] == "Acme Drywall Supply"


def test_api_daemon_status():
    res = client.get("/api/intake/daemon/status")
    assert res.status_code == 200
    data = res.json()
    assert "is_running" in data
    assert "interval_seconds" in data


def test_api_export_excel():
    res = client.get("/api/export/excel?status=committed")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "attachment" in res.headers["content-disposition"]


def test_api_spreadsheet_sync(tmp_path):
    target_xlsx = str(tmp_path / "Job101_PL.xlsx")
    sync_payload = {
        "job_name": "101 Main St Remodel",
        "spreadsheet_path": target_xlsx,
    }
    res = client.post("/api/spreadsheet/sync", json=sync_payload)
    assert res.status_code == 200
    data = res.json()
    assert data["ok"] is True
    assert data["job_name"] == "101 Main St Remodel"
    assert Path(target_xlsx).exists()


def test_api_check_email_not_configured(monkeypatch):
    """When IMAP credentials are not set, /api/intake/check-email returns a clear 'not configured' response without crashing."""
    monkeypatch.delenv("IMAP_HOST", raising=False)
    monkeypatch.delenv("IMAP_SERVER", raising=False)
    monkeypatch.delenv("IMAP_USERNAME", raising=False)
    monkeypatch.delenv("IMAP_USER", raising=False)
    monkeypatch.delenv("IMAP_PASSWORD", raising=False)

    res = client.get("/api/intake/check-email")
    assert res.status_code == 200
    data = res.json()
    assert data["configured"] is False
    assert "not configured" in data["status"]
    assert "not configured" in data["message"].lower()

    # Also test POST method
    res_post = client.post("/api/intake/check-email")
    assert res_post.status_code == 200
    assert res_post.json()["configured"] is False


def test_api_check_email_configured_poll_cycle(monkeypatch, tmp_path):
    """When IMAP is configured, /api/intake/check-email runs an on-demand poll cycle and ingests attachments."""
    import email.message
    from unittest.mock import patch, MagicMock

    monkeypatch.setenv("IMAP_HOST", "imap.contractor-mail.com")
    monkeypatch.setenv("IMAP_USERNAME", "inbox@contractor.com")
    monkeypatch.setenv("IMAP_PASSWORD", "supersecret")

    msg = email.message.EmailMessage()
    msg["Subject"] = "Invoice 84 Lumber"
    msg["From"] = "billing@84lumber.com"
    pdf_bytes = (Path(__file__).parent.parent / "sample_invoices" / "84_lumber_inv_301.pdf").read_bytes()
    msg.add_attachment(pdf_bytes, maintype="application", subtype="pdf", filename="84L_301.pdf")

    mock_imap = MagicMock()
    mock_imap.login.return_value = ("OK", [b"Logged in"])
    mock_imap.select.return_value = ("OK", [b"1"])
    mock_imap.search.return_value = ("OK", [b"101"])
    mock_imap.fetch.return_value = ("OK", [(b"101 (RFC822 {100}", msg.as_bytes()), b")"])
    mock_imap.store.return_value = ("OK", [b"Flags updated"])

    with patch("imaplib.IMAP4_SSL", return_value=mock_imap):
        res = client.get("/api/intake/check-email")
        assert res.status_code == 200
        data = res.json()
        assert data["configured"] is True
        assert data["status"] == "success"
        assert data["processed_count"] >= 1
        assert mock_imap.store.called


def test_api_job_export_xlsx(tmp_path):
    """GET /api/jobs/{job_id}/export.xlsx downloads a real, openpyxl-parseable Job P&L workbook."""
    import io
    import openpyxl

    res = client.get("/api/jobs/job-101/export.xlsx")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "101_Main_St_Remodel_PL.xlsx" in res.headers["content-disposition"]

    # Verify content is a valid, parseable Excel workbook
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Job P&L Summary" in wb.sheetnames
    assert "Cost Code Breakdown" in wb.sheetnames
    assert "Invoice Ledger" in wb.sheetnames

    ws_summary = wb["Job P&L Summary"]
    assert ws_summary["B4"].value == "101 Main St Remodel"
    assert ws_summary["D4"].value == 85000.0


def test_api_upload_image_triggers_vision_extractor(tmp_path, monkeypatch):
    """Uploading an image file to /api/upload triggers the Kimi-K3 vision extraction path."""
    import json
    from unittest.mock import patch, MagicMock

    monkeypatch.setenv("OPENAI_API_BASE", "https://api.openai.com/v1")
    monkeypatch.setenv("OPENAI_API_KEY", "mock-key-12345")

    fake_vision_response = {
        "choices": [
            {
                "message": {
                    "content": json.dumps({
                        "vendor": "White Cap",
                        "invoice_number": "WC-77102",
                        "invoice_date": "2026-08-18",
                        "po_number": "PO-101-Concrete",
                        "total_amount": 2350.00,
                        "tax_amount": 150.00,
                        "subtotal_amount": 2200.00,
                    })
                }
            }
        ]
    }

    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = fake_vision_response
    mock_resp.raise_for_status = MagicMock()

    with patch("requests.post", return_value=mock_resp) as mock_post:
        fake_image_bytes = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDRmockimage"
        res = client.post(
            "/api/upload",
            files={"files": ("receipt.png", fake_image_bytes, "image/png")}
        )
        assert res.status_code == 200
        data = res.json()
        assert data["processed"] == 1
        inv = data["results"][0]["invoice"]
        assert inv["vendor_name"] == "White Cap"
        assert inv["invoice_number"] == "WC-77102"
        assert inv["total_amount"] == 2350.00
        # Assert vision HTTP call was invoked
        assert mock_post.called
        call_json = mock_post.call_args[1]["json"]
        assert call_json["model"] == "moonshotai/Kimi-K3"

